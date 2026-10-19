"""報告資料（PPT）を作る。run_study.py・run_sensitivity.py の結果から作図し、社内報告テンプレートに流し込む。

使い方: python scripts/make_report.py [--template docs/report_template.pptx] [--out data/output/study/報告_モビリティシェア検討.pptx]
グラフは PowerPoint のネイティブグラフ（編集可能）。地図だけ画像。
"""
from __future__ import annotations

import argparse
import copy
import json
import datetime as dt
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from src.loaders import load_config, load_network, load_survey, resolve, survey_path  # noqa: E402
from src.plan import demand_hourly, fleet_for, make_ops, optimal_layout  # noqa: E402
from src.study import annual_cost, month_eval  # noqa: E402

_add_series = CategoryChartData.add_series


def _add_series_nan(self, name, values=(), number_format=None):
    """NaN（データ無し）は空欄にする。"""
    vals = [None if v is None or (isinstance(v, float) and np.isnan(v)) else v for v in values]
    return _add_series(self, name, vals, number_format)


CategoryChartData.add_series = _add_series_nan

# ---------------------------------------------------------------- 体裁（テンプレートの色・書体）
NAVY, TEAL, INK, MUTED, MID, LIGHT, BG, PURPLE = ("1A2F81", "086B89", "010817", "566A95", "808FB0", "ABB5CA",
                                                  "F7F8F9", "5D5491")
WHITE = "FFFFFF"
JP, JPB, JPM = "たづがね角ゴシック Info", "たづがね角ゴシック Info Bold", "たづがね角ゴシック Info Medium"
EB_C, KB_C = NAVY, TEAL
PAT_NAME = {"a": "a：夜1回", "b": "b：a＋9:30", "c": "c：b＋15:30"}


def rgb(h):
    return RGBColor.from_string(h)


def yen_man(x):
    return f"{x / 1e4:,.0f}万円"


# ---------------------------------------------------------------- スライド操作
def dup(prs, src):
    """テンプレートのスライドを複製（画像などの関係も付け替える）。"""
    new = prs.slides.add_slide(src.slide_layout)
    for shp in list(new.shapes):
        shp._element.getparent().remove(shp._element)
    rid = {}
    for r in src.part.rels.values():
        if "notesSlide" in r.reltype or "slideLayout" in r.reltype:
            continue
        rid[r.rId] = new.part.relate_to(r._target, r.reltype)
    bg = src._element.cSld.bg
    if bg is not None:
        new._element.cSld.insert(0, copy.deepcopy(bg))
    tree = new.shapes._spTree
    for el in src.shapes._spTree.iterchildren():
        if el.tag.endswith("}nvGrpSpPr") or el.tag.endswith("}grpSpPr"):
            continue
        e = copy.deepcopy(el)
        for node in e.iter():
            for k, v in list(node.attrib.items()):
                if k.startswith("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}") and v in rid:
                    node.set(k, rid[v])
        tree.append(e)
    return new


def by_id(slide, sid):
    for s in slide.shapes:
        if s.shape_id == sid:
            return s
    raise KeyError(sid)


def drop(slide, *ids):
    for i in ids:
        e = by_id(slide, i)._element
        e.getparent().remove(e)


def keep_only(slide, ids):
    for s in list(slide.shapes):
        if s.shape_id not in ids:
            s._element.getparent().remove(s._element)


def put(shape, lines):
    """文字の書式を保ったまま置き換える。lines：文字列 or 文字列のリスト（段落ごと）。"""
    if isinstance(lines, str):
        lines = [lines]
    tf = shape.text_frame
    ps = tf.paragraphs
    proto = copy.deepcopy(ps[0]._p)
    for p in ps[1:]:
        p._p.getparent().remove(p._p)
    first = ps[0]._p
    parent = first.getparent()
    parent.remove(first)
    for line in lines:
        p = copy.deepcopy(proto)
        runs = p.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}r")
        for r in runs[1:]:
            p.remove(r)
        for br in p.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}br"):
            p.remove(br)
        if runs:
            runs[0].find("{http://schemas.openxmlformats.org/drawingml/2006/main}t").text = line
        parent.append(p)


def tb(slide, x, y, w, h, paras, size=16, color=INK, font=JP, bold=False, align=PP_ALIGN.LEFT,
       anchor=MSO_ANCHOR.TOP, space_after=6, fill=None):
    """テキストボックス。paras：文字列 or (文字列, {size,color,bold,font}) のリスト。"""
    s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
        for m in ("margin_left", "margin_right"):
            setattr(tf, m, Inches(0.25))
        for m in ("margin_top", "margin_bottom"):
            setattr(tf, m, Inches(0.15))
    if isinstance(paras, (str, tuple)):
        paras = [paras]
    for i, item in enumerate(paras):
        text, o = (item, {}) if isinstance(item, str) else item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(o.get("after", space_after))
        r = p.add_run()
        r.text = text
        f = r.font
        f.size = Pt(o.get("size", size))
        f.bold = o.get("bold", bold)
        f.color.rgb = rgb(o.get("color", color))
        fn = o.get("font", JPB if o.get("bold", bold) else font)
        f.name = fn
        rPr = r._r.get_or_add_rPr()
        ea = rPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
        if ea is None:
            ea = etree.SubElement(rPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
        ea.set("typeface", fn)
    return s


def card(slide, x, y, w, h, fill=WHITE, line=LIGHT):
    from pptx.enum.shapes import MSO_SHAPE
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = 0.06
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    s.line.color.rgb = rgb(line)
    s.line.width = Pt(1)
    s.shadow.inherit = False
    return s


def table(slide, x, y, w, df, col_w=None, size=12, row_h=0.36, head_fill=NAVY, highlight=None, first_bold=True,
          left_cols=(0,)):
    nr, nc = df.shape[0] + 1, df.shape[1]
    gs = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    t = gs.table
    tblPr = t._tbl.tblPr
    for a in ("bandRow", "firstRow"):
        tblPr.set(a, "0")
    sid = tblPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId")
    if sid is not None:
        sid.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # 罫線のみ
    cw = col_w or [w / nc] * nc
    for j, c in enumerate(cw):
        t.columns[j].width = Inches(c)
    for i in range(nr):
        t.rows[i].height = Inches(row_h)
        for j in range(nc):
            cell = t.cell(i, j)
            v = df.columns[j] if i == 0 else df.iat[i - 1, j]
            cell.text = "" if v is None or (isinstance(v, float) and np.isnan(v)) else str(v)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j in left_cols else PP_ALIGN.CENTER
            for r in p.runs:
                r.font.size = Pt(size)
                head = i == 0
                b = head or (first_bold and j == 0)
                r.font.bold = b
                r.font.name = JPB if b else JP
                r.font.color.rgb = rgb(WHITE if head else INK)
            cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = rgb(head_fill)
            elif highlight is not None and highlight(i - 1):
                cell.fill.fore_color.rgb = rgb("E3E8F4")
            else:
                cell.fill.fore_color.rgb = rgb(WHITE if i % 2 else BG)
    return gs


def style_chart(ch, size=12, legend=True, legend_pos=XL_LEGEND_POSITION.TOP):
    ch.font.size = Pt(size)
    ch.font.name = JP
    ch.font.color.rgb = rgb(INK)
    ch.has_legend = legend
    if legend:
        ch.legend.position = legend_pos
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(size)
    try:
        va = ch.value_axis
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = rgb("E3E6EC")
        va.format.line.fill.background()
        va.tick_labels.font.size = Pt(size)
        va.tick_labels.font.color.rgb = rgb(MUTED)
        ca = ch.category_axis
        ca.format.line.color.rgb = rgb(LIGHT)
        ca.tick_labels.font.size = Pt(size)
        ca.has_major_gridlines = False
    except (ValueError, AttributeError):
        pass


def zero_axis(ch):
    ch.value_axis.minimum_scale = 0


def color_series(ch, colors, invisible=()):
    for i, s in enumerate(ch.plots[0].series):
        if i in invisible:
            s.format.fill.background()
            s.format.line.fill.background()
            continue
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = rgb(colors[i % len(colors)])


def chart_title(ch, text, size=14):
    ch.has_title = True
    tf = ch.chart_title.text_frame
    tf.text = text
    r = tf.paragraphs[0].runs[0]
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.name = JPB
    r.font.color.rgb = rgb(NAVY)


def box_chart(slide, x, y, w, h, cats, stats, title, fmt="0", colors=(LIGHT, NAVY, TEAL, LIGHT), size=12):
    """箱ひげ（積み上げ棒で表現）：最小（透明）・最小→25%・25%→中央値・中央値→75%・75%→最大。"""
    cd = CategoryChartData()
    cd.categories = cats
    st = np.array(stats, dtype=float)   # (n, 5) min q1 med q3 max
    segs = [st[:, 0], st[:, 1] - st[:, 0], st[:, 2] - st[:, 1], st[:, 3] - st[:, 2], st[:, 4] - st[:, 3]]
    names = ["最小", "最小〜25%", "25%〜中央値", "中央値〜75%", "75%〜最大"]
    for n_, s_ in zip(names, segs):
        cd.add_series(n_, [float(v) for v in s_], number_format=fmt)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    style_chart(ch, size=size, legend=False)
    color_series(ch, [WHITE, colors[0], colors[1], colors[2], colors[3]], invisible=(0,))
    ch.plots[0].gap_width = 80
    ch.plots[0].overlap = 100
    chart_title(ch, title)
    ch.value_axis.tick_labels.number_format = fmt
    ch.value_axis.tick_labels.number_format_is_linked = False
    return ch


# ---------------------------------------------------------------- データ
def q5(s):
    s = pd.Series(s).dropna()
    if not len(s):
        return [np.nan] * 5
    return [s.min(), s.quantile(.25), s.median(), s.quantile(.75), s.max()]


def rng_txt(s, fmt="{:.0f}"):
    s = pd.Series(s).dropna()
    if not len(s):
        return "未達"
    med = fmt.format(s.median())
    return med if s.min() == s.max() else f"{med}（{fmt.format(s.min())}–{fmt.format(s.max())}）"


def short_scen(s):
    return s.replace("（暫定）", "").replace("旧シナリオ", "旧")


def load_results(out):
    d = pd.read_csv(out / "results.csv")
    d["達成"] = d["達成"].astype(bool)
    return d


# ---------------------------------------------------------------- 地図（画像）
def port_map_png(net, cfg, ports_idx, assign, path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    for f in font_manager.findSystemFonts():
        if "ipag" in f.lower() or "japanese-gothic" in f.lower():
            font_manager.fontManager.addfont(f)
            plt.rcParams["font.family"] = font_manager.FontProperties(fname=f).get_name()
            break
    fig, ax = plt.subplots(figsize=(7.2, 6.2), dpi=200)
    nd = net.nodes
    ax.scatter(nd["lon"], nd["lat"], s=0.6, c="#D5DAE5", zorder=1)
    cols = ["#1A2F81", "#086B89", "#5D5491", "#808FB0", "#2E7D32"]
    h = net.houses
    ax.scatter(h["lon"], h["lat"], s=9, c=[cols[a % len(cols)] for a in assign], alpha=.55, zorder=2, lw=0)
    c = net.candidates.iloc[ports_idx]
    for i, (_, r) in enumerate(c.iterrows()):
        ax.scatter(r["lon"], r["lat"], s=260, marker="s", c=cols[i % len(cols)], edgecolors="white", lw=1.5, zorder=4)
        ax.annotate(f"P{i + 1}", (r["lon"], r["lat"]), color="white", ha="center", va="center", fontsize=8,
                    weight="bold", zorder=5)
    for key, mk in (("station", "D"), ("aeon", "D")):
        p = cfg["destinations"][key]["port"]
        ax.scatter(p["lon"], p["lat"], s=200, marker=mk, c="#E07A1F", edgecolors="white", lw=1.5, zorder=4)
        ax.annotate(cfg["destinations"][key]["port_label"], (p["lon"], p["lat"]), xytext=(8, 6),
                    textcoords="offset points", fontsize=9, color="#010817", weight="bold", zorder=5)
    lat = np.r_[h["lat"], [cfg["destinations"][k]["port"]["lat"] for k in ("station", "aeon")]]
    lon = np.r_[h["lon"], [cfg["destinations"][k]["port"]["lon"] for k in ("station", "aeon")]]
    pad = 0.0012
    ax.set_xlim(lon.min() - pad, lon.max() + pad * 2.5)
    ax.set_ylim(lat.min() - pad, lat.max() + pad)
    ax.set_aspect(1 / np.cos(np.radians(lat.mean())))
    ax.set_xticks([])
    ax.set_yticks([])
    for s in ax.spines.values():
        s.set_color("#ABB5CA")
    fig.tight_layout()
    fig.savefig(path, dpi=200)
    plt.close(fig)


# ---------------------------------------------------------------- 本体
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", default=str(ROOT / "docs" / "report_template.pptx"))
    ap.add_argument("--out", default=None)
    a = ap.parse_args()

    cfg = load_config()
    out_dir = resolve(cfg["paths"].get("output_dir", "data/output")) / "study"
    d = load_results(out_dir)
    sens = pd.read_csv(out_dir / "sensitivity.csv") if (out_dir / "sensitivity.csv").exists() else None
    buf = pd.read_csv(out_dir / "rebalance_buffer_test.csv") if (out_dir / "rebalance_buffer_test.csv").exists() else None
    meta_p = out_dir / "meta.json"
    districts = json.loads(meta_p.read_text(encoding="utf-8")).get("districts") if meta_p.exists() else None
    net = load_network(cfg).subset(districts)
    n_h = len(net.houses)
    dist_txt = "・".join(f"{k}{v}戸" for k, v in net.houses["街区"].value_counts(sort=False).items()) \
        if "街区" in net.houses else f"{n_h}戸"
    cfg["simulation"]["n_households"] = cfg["simulation"].get("n_households") or len(net.houses)
    sv = load_survey(survey_path(cfg)[0], cfg)
    P = cfg["survey_xlsx"]["default_modes"]
    scens = list(cfg["demand_scenarios"])
    caps = list(cfg["station_port"]["capacity_cases"].values())
    base_cap = cfg["station_port"]["max_slots"]
    reps = int(d["回"].nunique())
    n_cond = d.groupby(["需要シナリオ", "代替乗車", "運用", "駅ポート容量"]).ngroups

    def sel(scen=None, sub="なし(i)", pat="a", cap=base_cap, tgt=0.05):
        m = pd.Series(True, index=d.index)
        for col, v in (("需要シナリオ", scen), ("代替乗車", sub), ("運用", pat), ("駅ポート容量", cap), ("目標失敗率", tgt)):
            if v is not None:
                m &= d[col] == v
        return d[m]

    # 駅で返却失敗が起きた＝駅ポートが満車になった（朝から満車の日は最大在庫が記録されない旧版の結果を補正）
    fix = d["月_返却失敗率"].fillna(0) > 0
    d.loc[fix, "月_駅ポート最大在庫"] = np.maximum(d.loc[fix, "月_駅ポート最大在庫"].fillna(0), d.loc[fix, "駅ポート容量"])

    # ---- 推奨案（目標：失敗率5%・代替乗車なし。運用 a/b/c × 駅ポート容量のうち年間費用が最小、10回すべて達成）
    rec = {}
    for scen in scens:
        best = None
        for cap in caps:
            for pat in ("a", "b", "c"):
                x = sel(scen, pat=pat, cap=cap)
                ok = x[x["達成"]]
                if len(ok) < len(x) or not len(x):
                    continue
                c = ok["費用_年間総費用"].median()
                if best is None or c < best[2] - 1:
                    best = (pat, cap, c)
        if best is None:
            rec[scen] = None
            continue
        x = sel(scen, pat=best[0], cap=best[1])
        rec[scen] = dict(pattern=best[0], cap=best[1], nb=int(round(x["自転車"].median())),
                         nk=int(round(x["キックボード"].median())),
                         tot=rng_txt(x["合計台数"]), cost=x["費用_年間総費用"].median(), st_peak=int(x["月_駅ポート最大在庫"].max()),
                         ful=x["月_充足率"].median(), per_trip=x["費用_1トリップあたり費用"].median(),
                         per_veh=x["費用_1台あたり年間費用"].median(), util=x["費用_稼働率(回/台/日)"].median(),
                         reb=x["費用_運営費に占める再配置費用"].median(), slots=int(x["費用_区画数"].median()),
                         rows=x)
    lo_s, hi_s = scens[0], scens[1]

    def med_fleet(scen, sub="なし(i)", pat="a", cap=base_cap, tgt=0.05):
        x = sel(scen, sub, pat, cap, tgt)
        ok = x[x["達成"]]
        return float(ok["合計台数"].median()) if len(ok) and len(ok) >= len(x) / 2 else np.nan

    r15 = [med_fleet(sc, tgt=0.01) / med_fleet(sc) for sc in scens if not np.isnan(med_fleet(sc, tgt=0.01))]
    miss1 = [short_scen(sc) for sc in scens if np.isnan(med_fleet(sc, tgt=0.01))]
    red = [1 - med_fleet(sc, "あり(ii)") / med_fleet(sc) for sc in scens if not np.isnan(med_fleet(sc))]

    # ---- 推奨案の配置・1か月（低位・高位）
    lay, month = {}, {}
    for scen in (lo_s, hi_s):
        conv = cfg["demand_scenarios"][scen]
        oh, rh, sh = demand_hourly(sv, cfg, conv, P, "平日")
        oh2, rh2, _ = demand_hourly(sv, cfg, conv, P, "休日")
        r = rec[scen]
        cap_s = r["cap"] if r else base_cap
        L, ports, assign = optimal_layout(net, cfg, sh, 3, cfg["fleet_mix"]["ebike_share"], station_cap=cap_s)
        lay[scen] = (L, ports, assign)
        if r:
            ops = make_ops(cfg, pattern=r["pattern"], substitution=False)
            F = fleet_for(L, r["nb"], r["nk"], oh, rh, ops)
            summ, daily = month_eval(L, F, cfg, {"平日": (oh, rh), "休日": (oh2, rh2)}, pattern=r["pattern"],
                                     substitution=False, seed=1)
            cost = annual_cost(cfg, F, L.k, daily, cap_s)
            r["init"] = F.init
            month[scen] = (summ, daily, cost)
    png = out_dir / "_map.png"
    port_map_png(net, cfg, lay[lo_s][1], lay[lo_s][2], png)

    # ================================================================ PPT
    prs = Presentation(a.template)
    T = list(prs.slides)
    n_tmpl = len(T)
    W_ = prs.slide_width

    def frame(title, en, heading):
        s = dup(prs, T[10])
        keep_only(s, {2, 5, 20, 21, 22})
        put(by_id(s, 21), title)
        put(by_id(s, 22), en)
        put(by_id(s, 20), heading)
        return s

    today = dt.date.today()
    date_txt = f"{today.year}年{today.month}月{today.day}日"

    # 1 表紙
    s = dup(prs, T[0])
    put(by_id(s, 20), "南栗橋BLP街区")
    put(by_id(s, 22), "モビリティシェア")
    put(by_id(s, 21), "検討報告")
    put(by_id(s, 23), "配置・運用・費用の検討")
    put(by_id(s, 26), f"全{n_cond}条件×{reps}回のシミュレーション")
    put(by_id(s, 25), "作成：検討チーム")
    put(by_id(s, 24), f"作成日：{date_txt}")
    by_id(s, 24).width = Inches(3.8)

    # 2 結論（3カード）
    s = dup(prs, T[4])
    put(by_id(s, 32), "結論：推奨案")
    put(by_id(s, 33), "Recommendation")
    put(by_id(s, 31), "失敗率5%以下を最も安く満たす構成")
    rl, rh_ = rec[lo_s], rec[hi_s]
    put(by_id(s, 40), "PORT")
    put(by_id(s, 41), "住宅側ポート3か所")
    put(by_id(s, 46), "p-median／車種比率で加重")
    L0 = lay[lo_s][0]
    put(by_id(s, 42), [f"徒歩＋走行＋施設側徒歩の合計が最小になる3地点（P1〜P3）。各ポートの担当戸数は "
                       f"{'・'.join(str(int(round(x * len(net.houses)))) for x in L0.share)} 戸。",
                       f"駅西口ポート・モビリティステーションは固定。"])
    put(by_id(s, 47), "FLEET")
    put(by_id(s, 48), "混在フリートの台数")
    put(by_id(s, 53), "自転車60.3%：キックボード39.7%")
    fl = lambda r: f"自転車{r['nb']}＋キックボード{r['nk']}＝{r['nb'] + r['nk']}台" if r else "目標未達"
    put(by_id(s, 49), [f"低位：{fl(rl)}", f"高位：{fl(rh_)}",
                       f"（{reps}回の中央値。駅ポートの朝の初期配置を含む）"])
    put(by_id(s, 24), "OPERATION")
    put(by_id(s, 25), "駅ポート・運用・費用")
    put(by_id(s, 30), "運用a（夜1回の再配置）")
    put(by_id(s, 26), [f"低位：駅{rl['cap']}台・年間{yen_man(rl['cost'])}" if rl else "低位：—",
                       f"高位：駅{rh_['cap']}台・年間{yen_man(rh_['cost'])}" if rh_ else "高位：—",
                       "日中の再配置（b・c）は台数・費用とも増える。"])

    # 3 推奨案の文章＋地図
    s = frame("推奨案の内容", "Recommended Plan", "住宅側ポート・台数・駅ポート・運用・費用")
    s.shapes.add_picture(str(png), Inches(11.9), Inches(2.95), height=Inches(7.6))
    y = 3.05
    for scen in (lo_s, hi_s):
        r = rec[scen]
        if not r:
            continue
        stp = int(r["init"][L0.k].sum())
        txt = (f"住宅側ポート3か所・電動アシスト自転車{r['nb']}台＋電動キックボード{r['nk']}台・"
               f"駅ポート{r['cap']}台分（最大在庫{r['st_peak']}台、朝の初期配置{stp}台）・運用は{PAT_NAME[r['pattern']]}、"
               f"年間費用は{yen_man(r['cost'])}（1トリップ{r['per_trip']:,.0f}円・1台{r['per_veh'] / 1e4:,.1f}万円/年）。"
               f"1か月の充足率は{r['ful']:.1%}。")
        card(s, 1.15, y, 10.3, 3.35, fill=WHITE)
        tb(s, 1.45, y + 0.25, 9.8, 0.5, (f"{short_scen(scen)}需要（転換率{cfg['demand_scenarios'][scen]:.0%}・暫定）",
                                         {"bold": True, "color": NAVY, "size": 20}))
        tb(s, 1.45, y + 0.95, 9.7, 2.3, txt, size=16, color=INK)
        y += 3.65
    tb(s, 11.9, 10.6, 7.0, 0.4, "■ P1〜P3：住宅側ポート（点の色＝利用するポート）　◆ 駅西口ポート・モビリティステーション",
       size=11, color=MUTED)

    # 3b 新規街区の追加による変化（171戸の結果があるとき）
    prev_p = out_dir.parent / "study_171" / "results.csv"
    if prev_p.exists() and n_h != 171:
        d0 = load_results(prev_p.parent)

        def med0(df, scen, tgt, col="合計台数", pat="a", cap=base_cap):
            x = df[(df["需要シナリオ"] == scen) & (df["代替乗車"] == "なし(i)") & (df["運用"] == pat)
                   & (df["駅ポート容量"] == cap) & (df["目標失敗率"] == tgt)]
            ok = x[x["達成"]]
            return float(ok[col].median()) if len(ok) and len(ok) >= len(x) / 2 else np.nan

        s = frame("新規街区の追加による変化", "Effect of the New District",
                  f"171戸 → {n_h}戸（{dist_txt}）、代替乗車なし・運用a・駅ポート{base_cap}台")
        cd = CategoryChartData()
        cd.categories = [short_scen(x) for x in scens]
        cd.add_series("171戸", [med0(d0, sc, 0.05) for sc in scens])
        cd.add_series(f"{n_h}戸", [med0(d, sc, 0.05) for sc in scens])
        ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1.15), Inches(3.0), Inches(8.2), Inches(7.2), cd).chart
        style_chart(ch)
        color_series(ch, [LIGHT, NAVY])
        ch.plots[0].has_data_labels = True
        ch.plots[0].data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        ch.plots[0].data_labels.font.size = Pt(12)
        chart_title(ch, "失敗率5%の最小台数（中央値、空欄＝未達）")
        zero_axis(ch)

        def fmt_n(v):
            return "未達" if np.isnan(v) else f"{v:.0f}台"

        def fmt_c(v):
            return "—" if np.isnan(v) else f"{v / 1e4:,.0f}万円"

        def ratio(a_, b_):
            return "—" if np.isnan(a_) or np.isnan(b_) else f"{b_ / a_ - 1:+.0%}"

        rows = []
        for sc in scens:
            f0, f1 = med0(d0, sc, 0.20), med0(d, sc, 0.20)
            g0, g1 = med0(d0, sc, 0.05), med0(d, sc, 0.05)
            c0, c1 = med0(d0, sc, 0.05, "費用_年間総費用"), med0(d, sc, 0.05, "費用_年間総費用")
            rows.append([short_scen(sc), f"{fmt_n(f0)}→{fmt_n(f1)}", f"{fmt_n(g0)}→{fmt_n(g1)}（{ratio(g0, g1)}）",
                         f"{fmt_c(c0)}→{fmt_c(c1)}"])
        table(s, 9.7, 3.1, 9.15, pd.DataFrame(rows, columns=["需要", "充足率80%", "失敗率5%", "年間費用（5%）"]),
              col_w=[1.4, 2.2, 3.05, 2.5], size=12, row_h=0.6)
        tb(s, 9.7, 6.4, 9.15, 3.8, [
            f"戸数は{n_h / 171 - 1:+.0%}。需要は戸数に比例して増えるが、台数の増え方は需要シナリオと駅ポートの空きで変わる。",
            "新規街区はモビリティステーションから離れているため、住宅側ポートの位置も見直される（推奨案のページの地図）。",
            ("新規街区の区画内道路は地図（OSM）に無く、外周道路まで直線で接続した暫定計算。", {"size": 13, "color": MUTED})],
           size=14, space_after=8)

    # 4 検討の概要と条件
    s = frame("検討の概要と条件", "Overview & Conditions", f"全{n_cond}条件 × 各{reps}回（乱数を変えて）を計算")
    rows = [("需要シナリオ", "低位・高位（暫定）／旧シナリオ転換率10%・30%"),
            ("代替乗車", "(i) 希望車種が無ければ使わない／(ii) 他車種で代替"),
            ("リバランス", "a 夜1回／b a＋9:30 駅の余剰を戻す／c b＋15:30 住宅側と駅をそろえる"),
            ("駅ポート容量", "18台（区画①のみ）・25台（基本）・32台（区画②も2列）"),
            ("目標", "貸出失敗率 20%（充足率80%）・5%・1% それぞれの最小台数"),
            ("1回の計算", f"①平日・晴れの1日×{40}回で最小台数を探索 → ②その台数で1か月（平日・休日・気象）→ ③年間費用"),
            ("前提", f"住宅{n_h}戸（{dist_txt}）・駅は片道・イオンは往復・キックボードは電池交換で4時間停止")]
    table(s, 1.15, 3.1, 17.7, pd.DataFrame(rows, columns=["項目", "内容"]), col_w=[3.2, 14.5], size=15, row_h=0.62, left_cols=(0, 1))
    tb(s, 1.15, 8.3, 17.7, 1.8, [
        ("ばらつき＝同じ条件で乱数（利用希望の発生・気象の並び）だけを変えた10回の差。条件の差と区別して示す。",
         {"size": 15}),
        ("需要（転換率）と気象倍率は暫定値。東京都資料の数値に差し替えて再計算する（scripts/run_study.py）。",
         {"size": 15, "color": MUTED})])

    # 5 目標別の最小台数（表＋積み上げ棒）
    s = frame("目標ごとの最小台数", "Minimum Fleet by Target",
              f"代替乗車なし(i)・運用a・駅ポート{base_cap}台（{reps}回の中央値、括弧は最小–最大）")
    tg = [0.20, 0.05, 0.01]
    rows = []
    for scen in scens:
        rr = [short_scen(scen)]
        for t in tg:
            x = sel(scen, tgt=t)
            ok = x[x["達成"]]
            if len(ok) < len(x) / 2:
                rr.append(f"未達（{len(ok)}/{len(x)}回）")
            else:
                rr.append(f"{ok['自転車'].median():.0f}＋{ok['キックボード'].median():.0f}＝{rng_txt(ok['合計台数'])}")
        rows.append(rr)
    table(s, 1.15, 3.1, 9.6, pd.DataFrame(rows, columns=["需要", "充足率80%", "失敗率5%", "失敗率1%"]),
          col_w=[1.8, 2.6, 2.6, 2.6], size=13, row_h=0.62)
    tb(s, 1.15, 6.4, 9.6, 3.6, [
        "表記：自転車＋キックボード＝合計（中央値）",
        f"失敗率1%は5%の約{min(r15):.1f}〜{max(r15):.1f}倍の台数が要る。"
        + (f"{'・'.join(miss1)}では1%を半数以上の回で達成できない（探索範囲内）。" if miss1 else ""),
        f"代替乗車あり(ii)にすると、失敗率5%の台数は{min(red):.0%}〜{max(red):.0%}少なくなる（混在フリートのページ）。",
        f"需要が多いほど駅ポート{base_cap}台が満車になりやすく、台数が急増する（駅ポート容量のページ）。"],
       size=14, space_after=8)
    cd = CategoryChartData()
    cats = []
    eb, kb = [], []
    for scen in scens:
        for t in (0.20, 0.05):
            x = sel(scen, tgt=t)
            ok = x[x["達成"]]
            cats.append(f"{short_scen(scen)}\n{'80%' if t == 0.2 else '5%'}")
            eb.append(float(ok["自転車"].median()) if len(ok) >= len(x) / 2 else 0)
            kb.append(float(ok["キックボード"].median()) if len(ok) >= len(x) / 2 else 0)
    cd.categories = cats
    cd.add_series("電動アシスト自転車", eb)
    cd.add_series("電動キックボード", kb)
    ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(11.2), Inches(3.0), Inches(7.7), Inches(7.2), cd).chart
    style_chart(ch)
    color_series(ch, [EB_C, KB_C])
    ch.plots[0].gap_width = 60
    ch.plots[0].has_data_labels = True
    dl = ch.plots[0].data_labels
    dl.position = XL_LABEL_POSITION.CENTER
    dl.font.size = Pt(11)
    dl.font.color.rgb = rgb(WHITE)
    chart_title(ch, "最小台数の内訳（中央値、0＝未達）")
    zero_axis(ch)

    # 6 ばらつき：最小台数（全条件）
    s = frame("結果のばらつき①：最小台数", "Variability: Fleet Size",
              f"同じ条件で{reps}回計算したときの幅（箱＝25〜75%、薄色＝最小・最大）")
    for col, (t, ttl) in enumerate(((0.05, "失敗率5%の最小台数（合計）"), (0.20, "充足率80%の最小台数（合計）"))):
        cats, stats = [], []
        for scen in scens:
            for sub in ("なし(i)", "あり(ii)"):
                x = sel(scen, sub=sub, tgt=t)
                ok = x[x["達成"]]
                if len(ok) < len(x) / 2:
                    continue
                cats.append(f"{short_scen(scen)}{'(i)' if sub.startswith('なし') else '(ii)'}")
                stats.append(q5(ok["合計台数"]))
        box_chart(s, 1.15 + col * 8.95, 3.0, 8.75, 6.3, cats, stats, ttl)
    tb(s, 1.15, 9.5, 17.7, 0.9, "運用a・駅ポート25台。(i) 代替乗車なし／(ii) あり。幅は需要の偶然の揺れによるもので、計画台数は上側（75%点〜最大）を見て決めると安全。",
       size=14, color=MUTED)

    # 7 ばらつき：1か月の結果（全72条件）
    s = frame("結果のばらつき②：1か月の運用", "Variability: One-Month Simulation",
              "失敗率5%の台数で1か月（平日・休日・気象）を回した結果（全条件・全回）")
    x5 = d[(d["目標失敗率"] == 0.05) & d["達成"]]
    specs = [("月_充足率", "1か月の充足率", "0%"), ("月_返却失敗率", "返却失敗率（満車）", "0.0%"),
             ("月_代替乗車率", "代替乗車率（(ii)のみ）", "0%"), ("費用_年間総費用", "年間総費用（万円）", "#,##0")]
    for i, (col, ttl, fmt) in enumerate(specs):
        cats, stats = [], []
        for scen in scens:
            xx = x5[x5["需要シナリオ"] == scen]
            if col == "月_代替乗車率":
                xx = xx[xx["代替乗車"] == "あり(ii)"]
            if not len(xx):
                continue
            v = xx[col] / (1e4 if col.startswith("費用_年間") else 1)
            cats.append(short_scen(scen))
            stats.append(q5(v))
        box_chart(s, 1.15 + i * 4.45, 3.0, 4.3, 6.4, cats, stats, ttl, fmt=fmt, size=11)
    tb(s, 1.15, 9.6, 17.7, 0.8, "箱ひげは「代替乗車・運用a/b/c・駅ポート容量」の全組み合わせ×10回を需要ごとにまとめたもの。条件ごとの数値は付録の一覧表。",
       size=14, color=MUTED)

    # 8 混在フリート：代替乗車の効果
    s = frame("混在フリートと代替乗車", "Mixed Fleet & Substitution", "希望車種が無いときに他の車種で代替できるか")
    cd = CategoryChartData()
    cd.categories = [short_scen(x) for x in scens]
    for sub in ("なし(i)", "あり(ii)"):
        vals = []
        for scen in scens:
            x = sel(scen, sub=sub)
            ok = x[x["達成"]]
            vals.append(float(ok["合計台数"].median()) if len(ok) >= len(x) / 2 else 0)
        cd.add_series(f"代替乗車{sub}", vals)
    ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1.15), Inches(3.0), Inches(9.5), Inches(7.2), cd).chart
    style_chart(ch)
    color_series(ch, [NAVY, MID])
    ch.plots[0].has_data_labels = True
    ch.plots[0].data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    ch.plots[0].data_labels.font.size = Pt(12)
    chart_title(ch, "失敗率5%の最小台数（中央値、0＝未達）")
    zero_axis(ch)
    sub_rate = x5[x5["代替乗車"] == "あり(ii)"]["月_代替乗車率"]
    card(s, 11.1, 3.0, 7.75, 3.3, fill=BG, line=BG)
    tb(s, 11.4, 3.2, 7.2, 0.5, ("代替乗車率（(ii)・1か月）", {"bold": True, "color": NAVY, "size": 18}))
    tb(s, 11.4, 3.9, 7.2, 1.2, (f"{sub_rate.median():.1%}", {"bold": True, "color": TEAL, "size": 48}))
    tb(s, 11.4, 5.2, 7.2, 1.0, f"範囲 {sub_rate.min():.1%}〜{sub_rate.max():.1%}（全条件・全回）", size=14, color=MUTED)
    tb(s, 11.1, 6.7, 7.75, 3.4, [
        "車種の比率は調査Q16（自転車60.3%・キックボード39.7%）。利用者ごとに希望車種を割り当て、往復は同じ車種を希望。",
        f"代替を認めると、失敗率5%の台数を{min(red):.0%}〜{max(red):.0%}減らせる（需要が多いほど効果大）。代替乗車率は数%程度。",
        "代替を前提にする場合は、利用規約・料金で他車種への乗り換えを案内する仕組みが要る。"], size=14, space_after=8)

    # 9 駅ポート容量
    s = frame("駅ポートの容量（18・25・32台）", "Station Port Capacity", "満車で返却できない「返却失敗」と必要台数への影響")
    cd = CategoryChartData()
    cd.categories = [short_scen(x) for x in scens]
    for cap in caps:
        vals = []
        for scen in scens:
            x = sel(scen, cap=cap)
            ok = x[x["達成"]]
            vals.append(float(ok["合計台数"].median()) if len(ok) >= len(x) / 2 else 0)
        cd.add_series(f"{cap}台", vals)
    ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1.15), Inches(3.0), Inches(8.7), Inches(6.2), cd).chart
    style_chart(ch)
    color_series(ch, [LIGHT, NAVY, TEAL])
    ch.plots[0].has_data_labels = True
    ch.plots[0].data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    ch.plots[0].data_labels.font.size = Pt(11)
    chart_title(ch, "失敗率5%の最小台数（(i)・運用a、0＝未達）")
    zero_axis(ch)
    rows = []
    for scen in scens:
        rr = [short_scen(scen)]
        for cap in caps:
            x = sel(scen, cap=cap)
            ok = x[x["達成"]]
            rr.append(f"{ok['月_返却失敗率'].median():.1%}／{ok['月_駅ポート最大在庫'].median():.0f}台" if len(ok) else "—")
        rows.append(rr)
    table(s, 10.2, 3.1, 8.65, pd.DataFrame(rows, columns=["需要"] + [f"容量{c}台" for c in caps]),
          col_w=[1.7, 2.3, 2.3, 2.35], size=12, row_h=0.6)
    cap_lines = []
    for scen in (lo_s, hi_s):
        f_ = {c: med_fleet(scen, cap=c) for c in caps}
        cost_ = {c: sel(scen, cap=c)[lambda z: z["達成"]]["費用_年間総費用"].median() for c in caps}
        cap_lines.append(f"{short_scen(scen)}：" + "・".join(
            f"{c}台→{'未達' if np.isnan(f_[c]) else f'{f_[c]:.0f}台/{cost_[c] / 1e4:,.0f}万円'}" for c in caps))
    tb(s, 10.2, 6.3, 8.65, 4.0, [
        ("表：返却失敗率／駅ポートの最大在庫（1か月、中央値）", {"size": 12, "color": MUTED}),
        "満車で返せない車両は区画外に置かれ、夜の回収まで使えない。そのため夕方の帰宅で不足し、台数が増える。",
        *cap_lines,
        "低位なら25台で足りる。高位では32台（区画②も2列）にすると台数・費用とも下がる。"], size=14, space_after=8)

    # 10 リバランス a/b/c
    s = frame("リバランス運用 a・b・c の比較", "Rebalancing Patterns", f"日中に駅の余剰を戻しても台数は減らない（駅ポート{base_cap}台）")
    rows = []
    for scen in (lo_s, hi_s):
        for pat in ("a", "b", "c"):
            x = sel(scen, pat=pat)
            ok = x[x["達成"]]
            if not len(ok):
                rows.append([short_scen(scen), PAT_NAME[pat], "未達", "", "", "", ""])
                continue
            rows.append([short_scen(scen), PAT_NAME[pat], rng_txt(ok["合計台数"]), f"{ok['月_作業回数/日'].median():.1f}",
                         f"{ok['月_移動台数/日'].median():.1f}", f"{ok['月_作業時間(分/日)'].median():.0f}",
                         yen_man(ok["費用_年間総費用"].median())])
    table(s, 1.15, 3.1, 11.0, pd.DataFrame(rows, columns=["需要", "運用", "最小台数", "作業回/日", "移動台/日", "作業分/日", "年間費用"]),
          col_w=[1.2, 2.0, 2.3, 1.3, 1.3, 1.3, 1.6], size=12, row_h=0.55, highlight=lambda i: rows[i][1].startswith("a"))
    tb(s, 12.5, 3.0, 6.35, 4.2, [
        ("なぜ減らないか", {"bold": True, "color": NAVY, "size": 18}),
        "駅は片道利用で、朝に駅へ集まった車両がそのまま夕方の帰宅に使われる。9:30に「夕方の期待需要＋1台」を超える分を住宅側へ戻すと、夕方の偶然の増加に足りず失敗が増える。",
        "15:30のcは、bの後では動かす車両がほぼ無く、bと同じ結果。"], size=14, space_after=8)
    if buf is not None:
        bb = buf.copy()
        bb["需要シナリオ"] = bb["需要シナリオ"].map(short_scen)
        bb = bb[["需要シナリオ", "運用", "駅に残す台数の基準", "合計台数"]]
        tb(s, 12.5, 7.15, 6.35, 0.4, ("追加検証：駅に残す台数を95%点にした場合（1回）", {"bold": True, "size": 13, "color": MUTED}))
        table(s, 12.5, 7.6, 6.35, bb.rename(columns={"需要シナリオ": "需要", "駅に残す台数の基準": "駅に残す基準", "合計台数": "台数"}),
              col_w=[0.9, 0.7, 3.75, 1.0], size=10, row_h=0.32)
    tb(s, 1.15, 3.1 + 0.55 * (len(rows) + 1) + 0.4, 11.0, 2.6, [
        "運用aを推奨。今回の全条件（需要4種×容量3種×代替2種）で、b・cがaより台数・費用を減らした条件は無い。",
        "軽トラック：キックボード8台・自転車4台/往復、作業30分＋3分/台＋10分/追加往復。"], size=14, space_after=8)

    # 11 1か月の推移（推奨案）
    s = frame("1か月シミュレーション（推奨案）", "One-Month Simulation",
              f"{cfg['month']['year']}年{cfg['month']['month']}月：平日・休日（土日祝）・気象（平年の出現日数から抽選）")
    for i, scen in enumerate((lo_s, hi_s)):
        if scen not in month:
            continue
        summ, daily, cost = month[scen]
        cd = CategoryChartData()
        cd.categories = [f"{int(str(r['日付'])[8:10])}" for _, r in daily.iterrows()]
        cd.add_series("利用希望（回/日）", [float(v) for v in daily["利用希望"]])
        cd.add_series("利用（回/日）", [float(v) for v in daily["利用"]])
        ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1.15 + i * 8.95), Inches(3.0), Inches(8.75),
                                Inches(5.3), cd).chart
        style_chart(ch, size=11)
        color_series(ch, [LIGHT, NAVY if i == 0 else TEAL])
        ch.plots[0].gap_width = 40
        ch.plots[0].overlap = 100
        r = rec[scen]
        chart_title(ch, f"{short_scen(scen)}：自転車{r['nb']}＋キックボード{r['nk']}台（運用{r['pattern']}）")
        hol = daily[daily["曜日区分"] == "休日"]
        rain = daily[daily["気象"].astype(str).str.contains("雨")]
        tb(s, 1.15 + i * 8.95, 8.45, 8.75, 1.9, [
            f"充足率 {summ['充足率']:.1%}／返却失敗 {summ['返却失敗率']:.1%}／電池交換 {summ['電池交換/日']:.1f}回/日",
            f"休日{len(hol)}日・雨{len(rain)}日。休日・雨の日は利用希望が少なく、平日晴れが台数を決める。",
            f"稼働率 {summ['稼働率(回/台/日)']:.2f}回/台/日"], size=13, space_after=4)

    # 12 費用
    s = frame("年間費用と他事例との比較", "Annual Cost & Benchmarks", "推奨案（失敗率5%・代替乗車なし・運用a、駅ポートは費用が最小の容量）")
    items = ["車両（償却）", "車両維持", "ポート（償却）", "用地", "人件費（リバランス）", "軽トラック"]
    cd = CategoryChartData()
    cd.categories = [short_scen(sc) for sc in scens if rec.get(sc)]
    for it in items:
        cd.add_series(it, [float(rec[sc]["rows"][f"費用_{it}"].median() / 1e4) for sc in scens if rec.get(sc)],
                      number_format="#,##0")
    ch = s.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED, Inches(1.15), Inches(3.0), Inches(9.3), Inches(5.0), cd).chart
    style_chart(ch, size=12)
    color_series(ch, [NAVY, MID, TEAL, LIGHT, PURPLE, "C9CFDC"])
    ch.plots[0].gap_width = 50
    ch.plots[0].has_data_labels = True
    ch.plots[0].data_labels.position = XL_LABEL_POSITION.CENTER
    ch.plots[0].data_labels.font.size = Pt(10)
    ch.plots[0].data_labels.font.color.rgb = rgb(WHITE)
    ch.plots[0].data_labels.number_format = "#,##0"
    ch.plots[0].data_labels.number_format_is_linked = False
    chart_title(ch, "年間費用の内訳（万円、中央値）")
    zero_axis(ch)
    rl = rec[lo_s]
    rh2 = rec[hi_s] or rl
    rows = [["1台あたり年間費用", f"{rl['per_veh'] / 1e4:.1f}／{rh2['per_veh'] / 1e4:.1f}万円", "12.8万円（中野区）"],
            ["運営費に占める再配置費用", f"{rl['reb']:.0%}／{rh2['reb']:.0%}", "3〜4割"],
            ["稼働率（回/台/日）", f"{rl['util']:.2f}／{rh2['util']:.2f}", "1.9〜3.1"],
            ["1トリップあたり費用", f"{rl['per_trip']:,.0f}／{rh2['per_trip']:,.0f}円", "—"]]
    table(s, 10.8, 3.1, 8.05, pd.DataFrame(rows, columns=["指標（低位／高位）", "本検討", "他事例"]), col_w=[3.3, 2.4, 2.35],
          size=13, row_h=0.6)
    tb(s, 10.8, 6.3, 8.05, 4.0, [
        ("ずれの理由", {"bold": True, "color": NAVY, "size": 18}),
        f"稼働率が他事例の1/{1.9 / rl['util']:.0f}〜1/{3.1 / rl['util']:.0f}：{n_h}戸の閉じた街区で、失敗率5%を守るために朝夕のピークに合わせた台数を持つため。日中は多くが待機。",
        "再配置の割合がやや高い：毎晩の初期配置への戻しを全日で行う前提。",
        "1台あたり費用は他事例とほぼ同じ水準。"], size=14, space_after=8)
    tb(s, 1.15, 8.3, 9.3, 2.0, [
        f"単価15万円/台・3年償却、維持1.5万円/台/年、ポート20万円/か所・5年、用地600円/区画/月、人件費1,500円/時、軽トラック500円/回、平日245日・休日120日。",
    ], size=12, color=MUTED)

    # 13〜15 感度分析
    if sens is not None and len(sens):
        sens = sens.copy()
        sens["達成"] = sens["達成"].astype(bool)
        sv_a = sens[(sens["運用"] == "a") & sens["達成"]]
        base_d = round((cfg["demand_scenarios"][lo_s] + cfg["demand_scenarios"][hi_s]) / 2, 4)
        base_row = sv_a[(sv_a["項目"] == "需要（転換率）") & (np.isclose(sv_a["値"], base_d))]
        base_fleet = float(base_row["合計台数"].iloc[0]) if len(base_row) else np.nan

        # 推奨（費用最小の運用）を各点で
        def best_at(item, val):
            x = sens[(sens["項目"] == item) & np.isclose(sens["値"], val) & sens["達成"]]
            if not len(x):
                return None
            return x.loc[x["年間総費用"].idxmin()]

        base_best = best_at("需要（転換率）", base_d)
        base_cost = float(base_best["年間総費用"]) if base_best is not None else np.nan
        labels = {"需要（転換率）": lambda v: f"{v:.0%}", "自転車の比率": lambda v: f"{v:.0%}",
                  "駅ポート容量": lambda v: f"{v:.0f}台", "手続き時間（秒）": lambda v: f"{v:.0f}秒",
                  "歩行速度（km/h）": lambda v: f"{v:.1f}", "車両単価": lambda v: f"×{v:g}",
                  "ポート初期費用": lambda v: f"×{v:g}", "軽トラック費": lambda v: f"×{v:g}"}

        def tornado(slide, x, y, w, h, metric, base, items, title, scale=1.0, fmt="0"):
            recs = []
            for it in items:
                pts = sens[(sens["項目"] == it)]
                if not len(pts):
                    continue
                lo_v, hi_v = pts["値"].min(), pts["値"].max()
                if metric == "合計台数":
                    get = lambda v: sv_a[(sv_a["項目"] == it) & np.isclose(sv_a["値"], v)]["合計台数"]
                    lo_m, hi_m = get(lo_v), get(hi_v)
                    lo_m = float(lo_m.iloc[0]) if len(lo_m) else np.nan
                    hi_m = float(hi_m.iloc[0]) if len(hi_m) else np.nan
                else:
                    bl, bh = best_at(it, lo_v), best_at(it, hi_v)
                    lo_m = float(bl["年間総費用"]) if bl is not None else np.nan
                    hi_m = float(bh["年間総費用"]) if bh is not None else np.nan
                recs.append((f"{it}\n{labels[it](lo_v)}〜{labels[it](hi_v)}", (lo_m - base) * scale, (hi_m - base) * scale))
            recs = [r for r in recs if not (np.isnan(r[1]) and np.isnan(r[2]))]
            recs.sort(key=lambda r: abs(np.nan_to_num(r[1])) + abs(np.nan_to_num(r[2])))
            cd = CategoryChartData()
            cd.categories = [r[0] for r in recs]
            cd.add_series("下限の値のとき", [float(np.nan_to_num(r[1])) for r in recs], number_format=fmt)
            cd.add_series("上限の値のとき", [float(np.nan_to_num(r[2])) for r in recs], number_format=fmt)
            ch = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
            style_chart(ch, size=11)
            color_series(ch, [MID, NAVY])
            ch.plots[0].overlap = 100
            ch.plots[0].gap_width = 40
            ch.plots[0].has_data_labels = True
            ch.plots[0].data_labels.font.size = Pt(10)
            ch.plots[0].data_labels.number_format = '+' + fmt + ';-' + fmt + ';0'
            ch.plots[0].data_labels.number_format_is_linked = False
            ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
            chart_title(ch, title)
            return recs

        s = frame("感度分析①：最小台数", "Sensitivity: Fleet Size",
                  f"基準（需要{base_d:.0%}・駅{base_cap}台ほか）からの増減（失敗率5%・運用a）")
        recs_f = tornado(s, 1.15, 3.0, 11.2, 7.3, "合計台数", base_fleet,
                ["需要（転換率）", "自転車の比率", "駅ポート容量", "手続き時間（秒）", "歩行速度（km/h）"],
                f"最小台数の増減（台）　基準 {base_fleet:.0f}台")
        tb(s, 12.8, 3.0, 6.05, 7.2, [
            ("読み方", {"bold": True, "color": NAVY, "size": 18}),
            "棒の長さ＝その項目を下限・上限に動かしたときの最小台数の変化。上ほど影響が大きい。",
            f"最も効くのは「{recs_f[-1][0].splitlines()[0]}」、次が「{recs_f[-2][0].splitlines()[0]}」。"
            "駅ポートが満車になると返却失敗で台数が急増する。",
            "需要は暫定値のため、東京都資料での確定が最優先。",
            "手続き時間・歩行速度・自転車比率の影響は小さい。",
            ("各点1回の計算のため、±2台程度は乱数の揺れを含む。", {"size": 12, "color": MUTED})], size=14, space_after=10)

        s = frame("感度分析②：年間費用", "Sensitivity: Annual Cost", "各点で費用が最小の運用を選んだときの年間総費用の増減")
        tornado(s, 1.15, 3.0, 11.2, 7.3, "年間総費用", base_cost,
                ["需要（転換率）", "自転車の比率", "駅ポート容量", "手続き時間（秒）", "歩行速度（km/h）", "車両単価", "ポート初期費用",
                 "軽トラック費"], f"年間総費用の増減（万円）　基準 {base_cost / 1e4:,.0f}万円", scale=1e-4)
        tb(s, 12.8, 3.0, 6.05, 7.2, [
            ("読み方", {"bold": True, "color": NAVY, "size": 18}),
            "費用は需要（＝台数）と車両単価で大きく変わる。",
            "ポート・軽トラックの単価を0.5〜2倍に動かしても、推奨する運用は変わらない（次ページ）。",
            f"車両単価±30%は、そのまま車両償却（費用の約4〜5割）に効く。"], size=14, space_after=10)

        # 境目
        s = frame("推奨案が変わる境目", "Decision Boundaries", "各項目を動かしたとき、推奨案（台数・運用）が変わる値")
        rows = []
        for it in labels:
            pts = sorted(sens[sens["項目"] == it]["値"].unique())
            seq = []
            for v in pts:
                b = best_at(it, v)
                seq.append((v, None if b is None else (b["運用"], int(b["合計台数"]), int(b["自転車"]), int(b["キックボード"]),
                                                         float(b["年間総費用"]))))
            txt = []
            prev = None
            for v, b in seq:
                if b is None:
                    txt.append(f"{labels[it](v)}：未達")
                    prev = None
                    continue
                key = (b[0], b[1])
                if prev is None or key != prev:
                    txt.append(f"{labels[it](v)}：運用{b[0]}・{b[2]}＋{b[3]}={b[1]}台")
                prev = key
            pats = {b[0] for _, b in seq if b}
            note = "運用は全範囲でa" if pats == {"a"} else "運用が変わる：" + "→".join(sorted(pats))
            rows.append([it, "／".join(txt) if len(txt) > 1 else txt[0] + "（全範囲で同じ）", note])
        table(s, 1.15, 3.1, 17.7, pd.DataFrame(rows, columns=["項目", "推奨案（変わる値ごと）", "運用"]),
              col_w=[2.8, 11.9, 3.0], size=11, row_h=0.72, left_cols=(0, 1, 2))
    else:
        s = frame("感度分析", "Sensitivity", "感度分析の結果がありません")
        tb(s, 1.15, 3.2, 17.7, 1.0, "python scripts/run_sensitivity.py を実行してから作り直してください。", size=16)

    # 16 課題（暫定値）
    s = dup(prs, T[8])
    put(by_id(s, 19), "前提の確度と課題")
    put(by_id(s, 20), "Assumptions to Confirm")
    put(by_id(s, 18), "結果を左右する暫定値を確定させる")
    put(by_id(s, 24), "需要（転換率）が暫定")
    put(by_id(s, 25), ["低位12%・高位24%は仮置き。台数・費用に最も効く項目。",
                       "東京都の協議会資料（環境から取得不可）の値に差し替えて再計算する。"])
    put(by_id(s, 29), "気象倍率・出現日数が暫定")
    put(by_id(s, 30), ["国の資料を参考にした仮定値（要確認）。1か月の結果に影響。",
                       "気象庁の平年値（久喜・鴻巣など近傍）での確認が必要。"])
    put(by_id(s, 34), "利用実績・稼働率の差")
    put(by_id(s, 35), [f"稼働率{rl['util']:.2f}回/台/日は他事例（1.9〜3.1）より低い。実績が無く検証できない。",
                       "実証運用で利用記録をとり、時間帯別需要と車種比率を更新する。"])

    # 17 今後
    s = dup(prs, T[9])
    put(by_id(s, 19), "今後の進め方")
    put(by_id(s, 20), "Next Steps")
    put(by_id(s, 18), "暫定値の確定 → 再計算 → 実証での検証")
    put(by_id(s, 29), "STEP 01")
    put(by_id(s, 28), "需要の確定")
    by_id(s, 28).top += Inches(0.15)
    put(by_id(s, 27), "東京都資料で転換率を確定し、全条件を再計算")
    put(by_id(s, 30), "config.yaml の demand_scenarios を差し替え、run_study.py・make_report.py を再実行。")
    put(by_id(s, 40), "STEP 02")
    put(by_id(s, 39), "設計の確定")
    by_id(s, 39).top += Inches(0.15)
    put(by_id(s, 38), "住宅側3か所・運用a、駅ポートは需要で25か32台")
    put(by_id(s, 41), f"住宅の追加（現在{n_h}戸）に合わせてポート数と台数を見直す。")
    put(by_id(s, 50), "STEP 03")
    put(by_id(s, 49), "実証で検証")
    by_id(s, 49).top += Inches(0.15)
    put(by_id(s, 48), "利用記録で需要・稼働率・電池交換を検証")
    put(by_id(s, 51), "記録をシミュレーションに戻し、台数と運用を更新する。")

    # 18 前提一覧
    s = frame("主な前提値と出典", "Assumptions", "詳細は docs/前提と出典.md、値はすべて config.yaml")
    rows = [("住宅", f"{n_h}戸（{dist_txt}）", "確定"),
            ("調査", "2023年アンケート 23人（目的・時間帯・車種希望Q16）", "確定"),
            ("走行速度", "自転車12.6 km/h・キックボード14.0 km/h（走行ログ実測）", "確定"),
            ("ポート座標", "駅西口ポート 36.113177,139.712318／モビリティステーション 36.111012,139.709442", "確定"),
            ("駅ポート", "1台1m×2m、基本25台（区画①18＋区画②7）", "確定"),
            ("需要（転換率）", "低位12%・高位24%、旧10%・30%", "暫定"),
            ("気象", "雨・猛暑・強風・寒さの倍率と月別出現日数", "要確認"),
            ("新規街区の道路", "区画内道路は未収録。外周道路まで直線で接続（中央値41m）", "暫定"),
            ("電池", "キックボード航続30km、20%未満で交換・4時間停止", "仮定"),
            ("費用", "車両15万円・3年、維持1.5万円/年、ポート20万円・5年ほか", "仮定")]
    table(s, 1.15, 3.1, 17.7, pd.DataFrame(rows, columns=["項目", "値", "状態"]), col_w=[3.0, 12.5, 2.2], size=14, left_cols=(0, 1),
          row_h=0.58, highlight=lambda i: rows[i][2] in ("暫定", "要確認"))

    # 19 まとめ
    s = dup(prs, T[10])
    put(by_id(s, 21), "まとめ")
    put(by_id(s, 22), "Summary")
    put(by_id(s, 20), "推奨案と次の一歩")
    by_id(s, 18).width = Inches(12.8)
    put(by_id(s, 18), [
        f"住宅側ポート3か所・運用a（夜1回の再配置）を推奨。失敗率5%で、低位需要は{fl(rl)}・駅ポート{rl['cap']}台、"
        f"高位需要は{fl(rh_)}・駅ポート{rh_['cap']}台。",
        "日中のリバランス（b・c）は台数・費用を増やすだけで効果が無い。駅ポートが満車になると返却失敗で台数が急増するため、"
        "需要が高位に近ければ区画②を2列（32台）にする。",
        f"{reps}回のばらつきは最小台数で数台程度。計画は上側の値で見込むと安全。",
        "需要（転換率）が結果を最も左右するため、東京都資料での確定と再計算を最優先で進める。"])
    put(by_id(s, 23), "モビリティポート検討")

    # 付録：全条件の一覧
    for scen in scens:
        s = frame(f"付録：全条件の結果（{short_scen(scen)}）", "Appendix: All Conditions",
                  f"{reps}回の中央値（最小–最大）。台数＝自転車＋キックボード合計")
        rows = []
        for sub in ("なし(i)", "あり(ii)"):
            for pat in ("a", "b", "c"):
                for cap in caps:
                    rr = [sub, pat, f"{cap}"]
                    for t in (0.20, 0.05, 0.01):
                        x = sel(scen, sub, pat, cap, t)
                        ok = x[x["達成"]]
                        rr.append(rng_txt(ok["合計台数"]) if len(ok) == len(x) else
                                  (f"{rng_txt(ok['合計台数'])}［{len(ok)}/{len(x)}回達成］" if len(ok) else "未達"))
                    x = sel(scen, sub, pat, cap, 0.05)
                    ok = x[x["達成"]]
                    rr += [rng_txt(ok["月_充足率"] * 100, "{:.1f}") + "%" if len(ok) else "—",
                           rng_txt(ok["月_返却失敗率"] * 100, "{:.1f}") + "%" if len(ok) else "—",
                           rng_txt(ok["費用_年間総費用"] / 1e4, "{:,.0f}") if len(ok) else "—"]
                    rows.append(rr)
        df = pd.DataFrame(rows, columns=["代替", "運用", "駅容量", "台数 充足80%", "台数 失敗5%", "台数 失敗1%",
                                         "月充足率(5%)", "返却失敗(5%)", "年間費用万円(5%)"])
        table(s, 1.15, 2.95, 17.7, df, col_w=[1.3, 0.8, 1.0, 2.2, 2.2, 2.6, 2.4, 2.4, 2.8], size=10, row_h=0.395,
              first_bold=False)

    # テンプレートの元スライドを削除
    sld = prs.slides._sldIdLst
    for sid in list(sld)[:n_tmpl]:
        prs.part.drop_rel(sid.rId)
        sld.remove(sid)
    out = Path(a.out) if a.out else out_dir / "報告_モビリティシェア検討.pptx"
    prs.save(out)
    print("保存:", out)


if __name__ == "__main__":
    main()
